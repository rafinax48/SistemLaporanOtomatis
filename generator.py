import os
import sys
import json
import base64
import re
import uuid
import subprocess
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image as PILImage
from dotenv import load_dotenv

# Base directory
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Load config
load_dotenv(os.path.join(BASE_DIR, ".env"))
load_dotenv(r"C:\Users\HERITAGE\Documents\Kuliah\Semester 5\Sistem Laporan\.env.local")

API_KEY = os.getenv("NINEROUTER_API_KEY", "sk-e318a1c54061784e-5vo9dy-c115e232")
BASE_URL = os.getenv("NINEROUTER_BASE_URL", "http://127.0.0.1:20128/v1")
MODEL_ID = os.getenv("NINEROUTER_MODEL_ID", "gemini/gemini-3.5-flash-lite")

TEMPLATE_PATH = os.path.join(BASE_DIR, "Laporan SA 1.docx")
OUTPUT_DIR = os.path.join(BASE_DIR, "public", "outputs")
UPLOAD_DIR = os.path.join(BASE_DIR, "public", "uploads")

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

def save_base64_image(image_data, prefix="img"):
    """Saves base64 or data URL image to the uploads folder and returns its path."""
    if not image_data or not isinstance(image_data, str):
        return None
    
    if os.path.exists(image_data):
        return image_data
    relative_path = os.path.join(BASE_DIR, image_data.lstrip("/\\"))
    if os.path.exists(relative_path):
        return relative_path

    ext = "png"
    if "data:image/" in image_data:
        match = re.search(r"data:image/(\w+);base64,(.+)", image_data)
        if match:
            ext = match.group(1).lower()
            if ext == "jpeg": ext = "jpg"
            image_data = match.group(2)
        else:
            image_data = image_data.split(",")[-1]
    
    try:
        raw_bytes = base64.b64decode(image_data)
        filename = f"{prefix}_{uuid.uuid4().hex[:8]}.{ext}"
        filepath = os.path.join(UPLOAD_DIR, filename)
        with open(filepath, "wb") as f:
            f.write(raw_bytes)
        return filepath
    except Exception as e:
        print(f"Warning: Failed to decode base64 image: {e}", file=sys.stderr)
        return None

def normalize_images(input_data, list_key, single_key, prefix="img"):
    """
    Accepts either a list of image items:
      [{'order': 1, 'image': '...', 'caption': '...', 'url': '...'}, ...]
    or a single image string, or a JSON string.
    Returns a sorted list of dicts:
      [{'order': 1, 'path': '/path/to/img', 'caption': '...'}, ...]
    """
    raw_val = input_data.get(list_key) or input_data.get(single_key)
    if not raw_val:
        return []
    
    if isinstance(raw_val, str):
        raw_val = raw_val.strip()
        if raw_val.startswith("[") and raw_val.endswith("]"):
            try:
                raw_val = json.loads(raw_val)
            except Exception:
                pass

    items = []
    if isinstance(raw_val, list):
        for idx, item in enumerate(raw_val):
            if isinstance(item, dict):
                img_data = item.get("image") or item.get("url") or item.get("path")
                if img_data:
                    saved_path = save_base64_image(img_data, f"{prefix}_{idx+1}")
                    if saved_path:
                        items.append({
                            "order": int(item.get("order", idx + 1)),
                            "path": saved_path,
                            "caption": item.get("caption", "").strip()
                        })
            elif isinstance(item, str) and item.strip():
                saved_path = save_base64_image(item, f"{prefix}_{idx+1}")
                if saved_path:
                    items.append({
                        "order": idx + 1,
                        "path": saved_path,
                        "caption": ""
                    })
    elif isinstance(raw_val, str) and raw_val.strip():
        saved_path = save_base64_image(raw_val, prefix)
        if saved_path:
            items.append({
                "order": 1,
                "path": saved_path,
                "caption": ""
            })

    items.sort(key=lambda x: int(x.get("order", 999)))
    return items

def normalize_codes(input_data, list_key, single_key, default_title="Program Utama"):
    """
    Accepts list of code items:
      [{'order': 1, 'title': '...', 'code': '...'}, ...]
    or single string, or JSON string.
    Returns a sorted list of dicts:
      [{'order': 1, 'title': '...', 'code': '...'}, ...]
    """
    raw_val = input_data.get(list_key) or input_data.get(single_key)
    if not raw_val:
        return []

    if isinstance(raw_val, str):
        raw_val = raw_val.strip()
        if raw_val.startswith("[") and raw_val.endswith("]"):
            try:
                raw_val = json.loads(raw_val)
            except Exception:
                pass

    items = []
    if isinstance(raw_val, list):
        for idx, item in enumerate(raw_val):
            if isinstance(item, dict):
                c = item.get("code") or ""
                if c.strip():
                    items.append({
                        "order": int(item.get("order", idx + 1)),
                        "title": item.get("title", f"Program {idx+1}").strip() or f"Program {idx+1}",
                        "code": c
                    })
            elif isinstance(item, str) and item.strip():
                items.append({
                    "order": idx + 1,
                    "title": f"Program {idx+1}",
                    "code": item
                })
    elif isinstance(raw_val, str) and raw_val.strip():
        items.append({
            "order": 1,
            "title": default_title,
            "code": raw_val
        })

    items.sort(key=lambda x: int(x.get("order", 999)))
    return items

def format_font(run, size=11, name="Calibri", bold=False, italic=False, caps=False):
    """Formats a run with Calibri font, 11pt default, and specified styling."""
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:cs'), name)
    run.bold = bold
    run.italic = italic
    if caps:
        run.font.all_caps = True

def insert_image_proportional(paragraph, image_path, max_w_in=5.2, max_h_in=6.0):
    """
    Inserts image into paragraph with proportional dimensions,
    ensuring it does NOT exceed max_w_in and max_h_in.
    """
    if not image_path or not os.path.exists(image_path):
        return None
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)

    w_in = max_w_in
    h_in = max_h_in
    try:
        with PILImage.open(image_path) as im:
            px_w, px_h = im.size
        aspect = px_h / px_w if px_w > 0 else 1.0
        if (max_w_in * aspect) > max_h_in:
            h_in = max_h_in
            w_in = max_h_in / aspect
        else:
            w_in = max_w_in
            h_in = max_w_in * aspect
    except Exception as e:
        print(f"Warning: PIL failed reading image dimensions: {e}", file=sys.stderr)
        w_in = max_w_in
        h_in = None

    run = paragraph.add_run()
    if h_in:
        run.add_picture(image_path, width=Inches(w_in), height=Inches(h_in))
    else:
        run.add_picture(image_path, width=Inches(w_in))
    return paragraph

def set_table_width_exact_cm(table, width_cm=11.5):
    """
    Strictly forces table width to exactly width_cm (default 11.5 cm),
    disables autofit, and aligns table to LEFT (mengikuti posisi teks paragraf).
    """
    width_dxa = int((width_cm / 2.54) * 1440)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(width_dxa))

    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.width = Inches(width_cm / 2.54)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))

def create_code_table(doc, code_text, width_cm=11.5):
    """
    Creates an independent 1x1 table grid for code snippet with exact 11.5 cm width,
    aligned to left, with Courier New 10pt font.
    """
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    set_table_width_exact_cm(table, width_cm)
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(10.0)
    # Ensure Word explicitly applies Courier New via rFonts
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), 'Courier New')
    rFonts.set(qn('w:hAnsi'), 'Courier New')
    rFonts.set(qn('w:cs'), 'Courier New')
    return table

def clean_step_line(line):
    """
    Ensures points use '-' strictly.
    Strips leading '1. ', '1) ', '• ', '* ', '- ' etc.
    """
    line = str(line).strip()
    if not line:
        return ""
    cleaned = re.sub(r'^(?:\d+[\.\)]\s*|[a-zA-Z][\.\)]\s*|[•\-\*]\s*)', '', line).strip()
    return f"- {cleaned}" if cleaned else ""

def clean_caption_text(cap):
    """
    Strips unwanted '#1', '#2', '# 1' tokens and normalizes whitespace.
    """
    if not cap:
        return ""
    c = re.sub(r'#\s*\d+', '', str(cap)).strip()
    c = re.sub(r'\s{2,}', ' ', c).strip()
    return c

def convert_docx_to_pdf(docx_path, pdf_path):
    """Converts a DOCX file to PDF using Windows Word COM Automation."""
    abs_docx = os.path.abspath(docx_path).replace("'", "''")
    abs_pdf = os.path.abspath(pdf_path).replace("'", "''")
    
    ps_cmd = (
        f"$word = New-Object -ComObject Word.Application; "
        f"$word.Visible = $false; "
        f"try {{ "
        f"  $doc = $word.Documents.Open('{abs_docx}'); "
        f"  $doc.SaveAs([ref]'{abs_pdf}', [ref]17); "
        f"  $doc.Close(); "
        f"}} finally {{ "
        f"  $word.Quit(); "
        f"}}"
    )
    
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps_cmd],
            capture_output=True,
            text=True,
            timeout=45
        )
        return os.path.exists(pdf_path)
    except Exception as e:
        print(f"Warning converting PDF via Word COM: {e}", file=sys.stderr)
        return False

def generate_report(user_input):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Template Word tidak ditemukan pada: {TEMPLATE_PATH}")

    doc = Document(TEMPLATE_PATH)
    
    # 1. Identitas Laporan & Mahasiswa
    nama_laporan = user_input.get('nama_laporan') or user_input.get('namaLaporan') or "Praktikum Strategi Algoritma"
    materi = user_input.get('materi') or "Laporan Praktikum"
    tanggal = user_input.get('tanggal') or datetime.now().strftime("%A %d %B %Y")
    tempat = user_input.get('tempat') or "Laboratorium Komputer"
    
    nama_mahasiswa = user_input.get('nama_mahasiswa') or user_input.get('nama') or "Rafi Satya Prayoga"
    nim = user_input.get('nim') or "2400018208"

    # Normalized ordered images & codes
    pretest_images = normalize_images(user_input, 'images_pretest', 'img_pretest', 'pretest')
    laprak_codes = normalize_codes(user_input, 'codes_laprak', 'kode_program', 'Program Utama')
    laprak_images = normalize_images(user_input, 'images_laprak', 'img_laprak', 'laprak')
    posttest_codes = normalize_codes(user_input, 'codes_posttest', 'kode_posttest', 'Program Posttest')
    posttest_images = normalize_images(user_input, 'images_posttest', 'img_posttest', 'posttest')

    # =========================================================================
    # A. COVER - Font Calibri, Ukuran 11 pt, Tanpa Duplikasi
    # =========================================================================
    # Paragraph 2: LAPORAN PRAKTIKUM
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = ""
        run2 = doc.paragraphs[2].add_run("LAPORAN PRAKTIKUM")
        format_font(run2, size=11, name="Calibri", bold=True)
        doc.paragraphs[2].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 3: Nama Praktikum
    if len(doc.paragraphs) > 3:
        doc.paragraphs[3].text = ""
        run3 = doc.paragraphs[3].add_run(nama_laporan.title())
        format_font(run3, size=11, name="Calibri", bold=True)
        doc.paragraphs[3].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 4: Materi
    if len(doc.paragraphs) > 4:
        doc.paragraphs[4].text = ""
        clean_materi = re.sub(r'^(?:praktikum\s*\d+\s*:\s*|minggu\s*\d+\s*:\s*)+', '', materi, flags=re.IGNORECASE).strip()
        run4 = doc.paragraphs[4].add_run(materi.upper())
        format_font(run4, size=11, name="Calibri", bold=True)
        doc.paragraphs[4].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 5: Jadwal / Tempat (Pencegahan Teks Duplikat)
    if len(doc.paragraphs) > 5:
        doc.paragraphs[5].text = ""
        clean_tgl = str(tanggal).strip()
        clean_tmp = str(tempat).strip()
        if clean_tmp.lower() in clean_tgl.lower():
            cover_jadwal = clean_tgl
        elif clean_tgl.lower() in clean_tmp.lower():
            cover_jadwal = clean_tmp
        else:
            cover_jadwal = f"{clean_tgl} / {clean_tmp}" if clean_tmp else clean_tgl
        
        run5 = doc.paragraphs[5].add_run(cover_jadwal)
        format_font(run5, size=11, name="Calibri", bold=True)
        doc.paragraphs[5].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 14: Disusun Oleh: (Calibri 11 Bold Center)
    if len(doc.paragraphs) > 14:
        doc.paragraphs[14].text = ""
        run14 = doc.paragraphs[14].add_run("Disusun Oleh:")
        format_font(run14, size=11, name="Calibri", bold=True)
        doc.paragraphs[14].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 15: Nama Mahasiswa (TIDAK BOLD, Calibri 11 Center)
    if len(doc.paragraphs) > 15:
        doc.paragraphs[15].text = ""
        run_nama = doc.paragraphs[15].add_run(nama_mahasiswa)
        format_font(run_nama, size=11, name="Calibri", bold=False)
        doc.paragraphs[15].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 16: NIM Mahasiswa Saja Tanpa Kelas (TIDAK BOLD, Calibri 11 Center)
    if len(doc.paragraphs) > 16:
        doc.paragraphs[16].text = ""
        nim_clean = str(nim).strip()
        run_nim = doc.paragraphs[16].add_run(nim_clean)
        format_font(run_nim, size=11, name="Calibri", bold=False)
        doc.paragraphs[16].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 21..25: Institusi UAD (Calibri 11 Bold Center)
    institusi_map = {
        21: "PROGRAM STUDI S1 INFORMATIKA",
        22: "FAKULTAS TEKNOLOGI INDUSTRI",
        23: "UNIVERSITAS AHMAD DAHLAN",
        25: "2026"
    }
    for idx_p, txt in institusi_map.items():
        if len(doc.paragraphs) > idx_p:
            doc.paragraphs[idx_p].text = ""
            r = doc.paragraphs[idx_p].add_run(txt)
            format_font(r, size=11, name="Calibri", bold=True)
            doc.paragraphs[idx_p].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # =========================================================================
    # B. PRETEST
    # - Judul "Pre Test": Calibri 11 Bold
    # - Pertanyaan: Calibri 11 TIDAK BOLD, HANYA Abjad Kapital (A, B, C...)
    # - "LAMPIRAN" dan "Lampiran A, B, dan C": TIDAK BOLD
    # - 1 Gambar = 1 Halaman Lampiran Rapi
    # =========================================================================
    if len(doc.paragraphs) > 27:
        doc.paragraphs[27].text = ""
        run_pt_title = doc.paragraphs[27].add_run("Pre Test")
        format_font(run_pt_title, size=11, name="Calibri", bold=True)
        doc.paragraphs[27].alignment = WD_ALIGN_PARAGRAPH.LEFT

    pretest_q_raw = user_input.get('pretest_q') or user_input.get('pretestQ') or ""
    raw_lines = [q.strip() for q in str(pretest_q_raw).split("\n") if q.strip()]
    if not raw_lines:
        raw_lines = [
            "Jelaskan tujuan dari praktikum ini.",
            "Sebutkan konsep dasar yang digunakan."
        ]

    # Format penomoran soal hanya menggunakan abjad kapital A., B., C.
    q_formatted = []
    for idx_q, q_item in enumerate(raw_lines):
        letter = chr(65 + idx_q) if idx_q < 26 else f"Q{idx_q+1}"
        cleaned = re.sub(r'^(?:[0-9]+[\.\)]\s*|[a-zA-Z][\.\)]\s*|[•\-\*]\s*)', '', q_item).strip()
        q_formatted.append(f"{letter}. {cleaned}")

    # Bersihkan placeholder lama pertanyaan di template (p29..p32)
    for i in [29, 30, 31, 32]:
        if len(doc.paragraphs) > i:
            doc.paragraphs[i].text = ""

    # Isi pertanyaan Pretest (Calibri 11 TIDAK BOLD)
    if len(doc.paragraphs) > 29 and q_formatted:
        doc.paragraphs[29].text = ""
        doc.paragraphs[29].paragraph_format.space_before = Pt(2)
        doc.paragraphs[29].paragraph_format.space_after = Pt(2)
        r_q0 = doc.paragraphs[29].add_run(q_formatted[0])
        format_font(r_q0, size=11, name="Calibri", bold=False)
        doc.paragraphs[29].alignment = WD_ALIGN_PARAGRAPH.LEFT

        for q_idx, q_text in enumerate(q_formatted[1:], start=1):
            target_slot = 29 + q_idx
            if target_slot < 33 and len(doc.paragraphs) > target_slot:
                doc.paragraphs[target_slot].text = ""
                doc.paragraphs[target_slot].paragraph_format.space_before = Pt(2)
                doc.paragraphs[target_slot].paragraph_format.space_after = Pt(2)
                r_qk = doc.paragraphs[target_slot].add_run(q_text)
                format_font(r_qk, size=11, name="Calibri", bold=False)
                doc.paragraphs[target_slot].alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p_extra_q = doc.paragraphs[33].insert_paragraph_before()
                p_extra_q.paragraph_format.space_before = Pt(2)
                p_extra_q.paragraph_format.space_after = Pt(2)
                r_qk = p_extra_q.add_run(q_text)
                format_font(r_qk, size=11, name="Calibri", bold=False)
                p_extra_q.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Header LAMPIRAN & Lampiran A, B, dan C (TIDAK BOLD, Calibri 11)
    if len(doc.paragraphs) > 33:
        doc.paragraphs[33].text = ""
        doc.paragraphs[33].paragraph_format.space_before = Pt(12)
        doc.paragraphs[33].paragraph_format.space_after = Pt(2)
        r_lamp = doc.paragraphs[33].add_run("LAMPIRAN")
        format_font(r_lamp, size=11, name="Calibri", bold=False)
        doc.paragraphs[33].alignment = WD_ALIGN_PARAGRAPH.CENTER

    if len(doc.paragraphs) > 34:
        doc.paragraphs[34].text = ""
        doc.paragraphs[34].paragraph_format.space_before = Pt(2)
        doc.paragraphs[34].paragraph_format.space_after = Pt(6)
        r_lamp_sub = doc.paragraphs[34].add_run("Lampiran A, B, dan C")
        format_font(r_lamp_sub, size=11, name="Calibri", bold=False, italic=True)
        doc.paragraphs[34].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Hapus seluruh tabel lama template dan seluruh paragraf legacy template (>= 35)
    # untuk mencegah teks duplikat, strip kosong '-', atau huruf melayang 'B.'
    for tbl in list(doc.tables):
        tbl._element.getparent().remove(tbl._element)

    while len(doc.paragraphs) > 35:
        p_rem = doc.paragraphs[-1]
        p_rem._element.getparent().remove(p_rem._element)

    # Masukkan Gambar Pretest (1 Gambar = 1 Halaman Lampiran Rapi)
    if pretest_images:
        for idx_pi, p_img in enumerate(pretest_images):
            if idx_pi > 0:
                p_pb_lamp = doc.add_paragraph()
                p_pb_lamp.paragraph_format.space_before = Pt(0)
                p_pb_lamp.paragraph_format.space_after = Pt(0)
                p_pb_lamp.add_run().add_break(WD_BREAK.PAGE)

            p_img_box = doc.add_paragraph()
            p_img_box.paragraph_format.keep_with_next = True
            insert_image_proportional(p_img_box, p_img["path"], max_w_in=5.2, max_h_in=6.6)

            p_cap_box = doc.add_paragraph()
            p_cap_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap_box.paragraph_format.space_before = Pt(2)
            p_cap_box.paragraph_format.space_after = Pt(6)
            clean_c = clean_caption_text(p_img["caption"])
            cap_text = clean_c if clean_c else f"Jawaban PreTest Lembar {idx_pi+1}"
            r_cap = p_cap_box.add_run(f"Gambar 1.{idx_pi+1}: {cap_text}")
            format_font(r_cap, size=10, name="Calibri", italic=True)

    # =========================================================================
    # C. LAPRAK (Hasil Praktikum)
    # - II. Hasil Praktikum (Calibri 11 Bold)
    # - Sub-bab wajib huruf abjad besar & BOLD:
    #   A. Alat dan Bahan
    #   B. Langkah Kerja
    #   C. Implementasi Kode Program (KODE DULU SEMUA, LEBAR TABEL 11.5 CM)
    #   D. Implementasi Screenshot (SEMUA GAMBAR MENYUSUL, CAPTION MENEMPEL)
    #   E. Analisis dan Ulasan
    # =========================================================================
    # Page Break ke Bagian Laprak
    pb_lap = doc.add_paragraph()
    pb_lap.paragraph_format.space_before = Pt(0)
    pb_lap.paragraph_format.space_after = Pt(0)
    pb_lap.add_run().add_break(WD_BREAK.PAGE)

    p_hp = doc.add_paragraph()
    p_hp.paragraph_format.space_before = Pt(6)
    p_hp.paragraph_format.space_after = Pt(4)
    r_hp = p_hp.add_run("II.  Hasil Praktikum")
    format_font(r_hp, size=11, name="Calibri", bold=True)

    # A. Alat dan Bahan (BOLD)
    p_ab_h = doc.add_paragraph()
    p_ab_h.paragraph_format.space_before = Pt(6)
    p_ab_h.paragraph_format.space_after = Pt(2)
    r_ab_h = p_ab_h.add_run("A. Alat dan Bahan")
    format_font(r_ab_h, size=11, name="Calibri", bold=True)

    laprak_bahan_raw = user_input.get('laprak_bahan') or user_input.get('laprakBahan')
    bahan_list = []
    if laprak_bahan_raw:
        if isinstance(laprak_bahan_raw, list):
            bahan_list = [str(b).strip() for b in laprak_bahan_raw if str(b).strip()]
        else:
            bahan_list = [b.strip() for b in str(laprak_bahan_raw).split(",") if b.strip()]
    if not bahan_list:
        bahan_list = ["PC / Laptop", "Compiler C++ (Dev-C++ / GCC)", "Library standar <iostream>"]

    for b_item in bahan_list:
        clean_b = re.sub(r'^(?:[0-9]+[\.\)]\s*|[a-zA-Z][\.\)]\s*|[•\-\*]\s*)', '', b_item).strip()
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(1.5)
        p_b.paragraph_format.line_spacing = 1.15
        r_b = p_b.add_run(f"- {clean_b}")
        format_font(r_b, size=11, name="Calibri")

    # B. Langkah Kerja (BOLD, Strict Dash Poin "-")
    p_lk_h = doc.add_paragraph()
    p_lk_h.paragraph_format.space_before = Pt(8)
    p_lk_h.paragraph_format.space_after = Pt(2)
    r_lk_h = p_lk_h.add_run("B. Langkah Kerja")
    format_font(r_lk_h, size=11, name="Calibri", bold=True)

    laprak_langkah_raw = user_input.get('laprak_langkah') or user_input.get('laprakLangkah')
    langkah_list = []
    if laprak_langkah_raw:
        if isinstance(laprak_langkah_raw, list):
            langkah_list = [clean_step_line(s) for s in laprak_langkah_raw if clean_step_line(s)]
        else:
            langkah_list = [clean_step_line(s) for s in str(laprak_langkah_raw).split("\n") if clean_step_line(s)]
    if not langkah_list:
        langkah_list = [
            "- Menyiapkan program C++ untuk menerima input data.",
            "- Melakukan kompilasi dan pengujian algoritma.",
            "- Mengamati dan mencatat hasil eksekusi program."
        ]

    for l_item in langkah_list:
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_before = Pt(1)
        p_l.paragraph_format.space_after = Pt(1.5)
        p_l.paragraph_format.line_spacing = 1.15
        r_l = p_l.add_run(l_item)
        format_font(r_l, size=11, name="Calibri")

    # C. Implementasi Kode Program (BOLD) - KODE DULU SEMUA
    p_kp_h = doc.add_paragraph()
    p_kp_h.paragraph_format.space_before = Pt(8)
    p_kp_h.paragraph_format.space_after = Pt(2)
    r_kp_h = p_kp_h.add_run("C. Implementasi Kode Program")
    format_font(r_kp_h, size=11, name="Calibri", bold=True)

    if laprak_codes:
        for c_idx, c_item in enumerate(laprak_codes, start=1):
            p_code_title = doc.add_paragraph()
            p_code_title.paragraph_format.space_before = Pt(6)
            p_code_title.paragraph_format.space_after = Pt(2)
            p_code_title.paragraph_format.keep_with_next = True
            r_ct = p_code_title.add_run(f"Tabel Kode 2.{c_idx}: {c_item.get('title', f'Program {c_idx}')}")
            format_font(r_ct, size=11, name="Calibri", bold=True)

            # Tabel Kode Lebar Konsisten Tepat 11.5 cm
            create_code_table(doc, c_item.get("code", ""), width_cm=11.5)

    # D. Implementasi Screenshot (BOLD) - SEMUA GAMBAR MENYUSUL
    p_is_h = doc.add_paragraph()
    p_is_h.paragraph_format.space_before = Pt(10)
    p_is_h.paragraph_format.space_after = Pt(2)
    r_is_h = p_is_h.add_run("D. Implementasi Screenshot")
    format_font(r_is_h, size=11, name="Calibri", bold=True)

    if laprak_images:
        for l_idx, l_img in enumerate(laprak_images, start=1):
            p_img_l = doc.add_paragraph()
            p_img_l.paragraph_format.keep_with_next = True
            insert_image_proportional(p_img_l, l_img["path"], max_w_in=5.2, max_h_in=4.5)

            p_cap_l = doc.add_paragraph()
            p_cap_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap_l.paragraph_format.space_before = Pt(2)
            p_cap_l.paragraph_format.space_after = Pt(6)
            clean_c = clean_caption_text(l_img.get("caption", ""))
            cap_l = clean_c if clean_c else f"Hasil Eksekusi Program {l_idx}"
            r_cap_l = p_cap_l.add_run(f"Gambar 2.{l_idx}: {cap_l}")
            format_font(r_cap_l, size=10, name="Calibri", italic=True)

    # E. Analisis dan Ulasan (BOLD)
    p_au_h = doc.add_paragraph()
    p_au_h.paragraph_format.space_before = Pt(10)
    p_au_h.paragraph_format.space_after = Pt(2)
    r_au_h = p_au_h.add_run("E. Analisis dan Ulasan")
    format_font(r_au_h, size=11, name="Calibri", bold=True)

    laprak_analisis_raw = user_input.get('laprak_analisis') or user_input.get('laprakAnalisis') or ""
    an_lines = [a.strip() for a in str(laprak_analisis_raw).split("\n") if a.strip()]
    if not an_lines:
        an_lines = [
            "Berdasarkan hasil pengujian yang dilakukan, program berhasil dieksekusi dengan baik sesuai dengan tujuan praktikum.",
            "Struktur logika program dan algoritma yang diimplementasikan menghasilkan output yang konsisten dan stabil."
        ]

    for a_item in an_lines:
        p_an = doc.add_paragraph()
        p_an.paragraph_format.space_before = Pt(2)
        p_an.paragraph_format.space_after = Pt(4)
        p_an.paragraph_format.line_spacing = 1.15
        p_an.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_an = p_an.add_run(a_item)
        format_font(r_an, size=11, name="Calibri")

    # =========================================================================
    # D. POSTTEST
    # - III. Post Test (Calibri 11 Bold)
    # - A. Tujuan Pembuatan Program (BOLD)
    # - B. Implementasi Kode Program (BOLD, TABEL 11.5 CM)
    # - C. Implementasi Screenshot (BOLD, GAMBAR TERPUSAT, CAPTION MENEMPEL)
    # =========================================================================
    pb_post = doc.add_paragraph()
    pb_post.paragraph_format.space_before = Pt(0)
    pb_post.paragraph_format.space_after = Pt(0)
    pb_post.add_run().add_break(WD_BREAK.PAGE)

    p_post_h = doc.add_paragraph()
    p_post_h.paragraph_format.space_before = Pt(6)
    p_post_h.paragraph_format.space_after = Pt(4)
    r_post_h = p_post_h.add_run("III. Post Test")
    format_font(r_post_h, size=11, name="Calibri", bold=True)

    # A. Tujuan Pembuatan Program (BOLD)
    p_pt_tj = doc.add_paragraph()
    p_pt_tj.paragraph_format.space_before = Pt(6)
    p_pt_tj.paragraph_format.space_after = Pt(2)
    r_pt_tj = p_pt_tj.add_run("A. Tujuan Pembuatan Program")
    format_font(r_pt_tj, size=11, name="Calibri", bold=True)

    posttest_tujuan_raw = user_input.get('posttest_tujuan') or user_input.get('posttestTujuan') or ""
    tj_lines = [t.strip() for t in str(posttest_tujuan_raw).split("\n") if t.strip()]
    if not tj_lines:
        tj_lines = [
            "Tujuan dari pembuatan program tugas mandiri / posttest ini adalah untuk memperdalam pemahaman mengenai materi praktikum yang telah dipelajari.",
            "Melalui program ini, mahasiswa dapat menguji secara mandiri penyelesaian studi kasus dengan implementasi algoritma yang efisien."
        ]

    for t_item in tj_lines:
        p_tj = doc.add_paragraph()
        p_tj.paragraph_format.space_before = Pt(2)
        p_tj.paragraph_format.space_after = Pt(4)
        p_tj.paragraph_format.line_spacing = 1.15
        p_tj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_tj = p_tj.add_run(t_item)
        format_font(r_tj, size=11, name="Calibri")

    # B. Implementasi Kode Program (BOLD) - Tabel 11.5 cm
    p_pt_kp = doc.add_paragraph()
    p_pt_kp.paragraph_format.space_before = Pt(8)
    p_pt_kp.paragraph_format.space_after = Pt(2)
    r_pt_kp = p_pt_kp.add_run("B. Implementasi Kode Program")
    format_font(r_pt_kp, size=11, name="Calibri", bold=True)

    if posttest_codes:
        for cp_idx, cp_item in enumerate(posttest_codes, start=1):
            p_cpt = doc.add_paragraph()
            p_cpt.paragraph_format.space_before = Pt(6)
            p_cpt.paragraph_format.space_after = Pt(2)
            p_cpt.paragraph_format.keep_with_next = True
            r_cpt = p_cpt.add_run(f"Tabel Kode 3.{cp_idx}: {cp_item.get('title', f'Program Posttest {cp_idx}')}")
            format_font(r_cpt, size=11, name="Calibri", bold=True)

            # Tabel Kode Lebar Konsisten Tepat 11.5 cm
            create_code_table(doc, cp_item.get("code", ""), width_cm=11.5)

    # C. Implementasi Screenshot (BOLD) - Gambar Terpusat & Menempel Caption
    p_pt_is = doc.add_paragraph()
    p_pt_is.paragraph_format.space_before = Pt(10)
    p_pt_is.paragraph_format.space_after = Pt(2)
    r_pt_is = p_pt_is.add_run("C. Implementasi Screenshot")
    format_font(r_pt_is, size=11, name="Calibri", bold=True)

    if posttest_images:
        for pi_idx, pi_img in enumerate(posttest_images, start=1):
            p_img_p = doc.add_paragraph()
            p_img_p.paragraph_format.keep_with_next = True
            insert_image_proportional(p_img_p, pi_img["path"], max_w_in=5.2, max_h_in=4.5)

            p_cap_p = doc.add_paragraph()
            p_cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap_p.paragraph_format.space_before = Pt(2)
            p_cap_p.paragraph_format.space_after = Pt(6)
            clean_c = clean_caption_text(pi_img.get("caption", ""))
            cap_p = clean_c if clean_c else f"Hasil Output PostTest {pi_idx}"
            r_cap_p = p_cap_p.add_run(f"Gambar 3.{pi_idx}: {cap_p}")
            format_font(r_cap_p, size=10, name="Calibri", italic=True)

    # =========================================================================
    # E. Output Penamaan File & Dokumen
    # =========================================================================
    custom_name = user_input.get('final_filename') or user_input.get('finalFilename')
    if custom_name:
        clean_name = re.sub(r'[^a-zA-Z0-9_-]', '_', os.path.splitext(custom_name)[0].strip())
    else:
        sanitized_materi = re.sub(r'[^a-zA-Z0-9_-]', '_', materi.strip())
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        clean_name = f"Laporan_{sanitized_materi}_{timestamp}"

    docx_filename = f"{clean_name}.docx"
    pdf_filename = f"{clean_name}.pdf"

    docx_path = os.path.join(OUTPUT_DIR, docx_filename)
    pdf_path = os.path.join(OUTPUT_DIR, pdf_filename)

    # Save DOCX
    doc.save(docx_path)

    # Convert to PDF via Word COM Automation
    pdf_created = convert_docx_to_pdf(docx_path, pdf_path)

    result = {
        "success": True,
        "filename": docx_filename,
        "docxPath": f"/outputs/{docx_filename}",
        "pdfPath": f"/outputs/{pdf_filename}" if pdf_created else None,
        "fullDocxPath": docx_path,
        "fullPdfPath": pdf_path if pdf_created else None
    }
    return result

if __name__ == "__main__":
    input_data = None
    
    if len(sys.argv) > 1:
        raw_arg = sys.argv[1].strip()
        if os.path.isfile(raw_arg):
            with open(raw_arg, "r", encoding="utf-8") as f:
                input_data = json.load(f)
        else:
            try:
                input_data = json.loads(raw_arg)
            except Exception:
                pass

    if not input_data and "GENERATOR_INPUT" in os.environ:
        try:
            input_data = json.loads(os.environ["GENERATOR_INPUT"])
        except Exception:
            pass

    if not input_data:
        input_data = {
            "nama_laporan": "Praktikum Strategi Algoritma",
            "materi": "PRAKTIKUM 1: KOMPLEKSITAS ALGORITMA",
            "tanggal": "Senin 13 April 2026 07.00-08.30",
            "tempat": "Laboratorium Basis Data",
            "nama_mahasiswa": "Rafi Satya Prayoga",
            "nim": "2400018208",
            "kelas": "C",
            "final_filename": "Laporan_Uji_Coba"
        }

    res = generate_report(input_data)
    print(json.dumps(res))
